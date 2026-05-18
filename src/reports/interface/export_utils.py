"""
报告导出工具
"""
from __future__ import annotations

import re
from html import escape
from pathlib import Path
from typing import Iterable

from django.conf import settings
from django.utils import timezone

from reports.models.report import Report

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None

try:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
except ImportError:  # pragma: no cover
    A4 = None
    colors = None
    getSampleStyleSheet = None
    HRFlowable = None
    mm = None
    ParagraphStyle = None
    pdfmetrics = None
    TA_CENTER = None
    TA_LEFT = None
    Table = None
    TableStyle = None
    UnicodeCIDFont = None
    Paragraph = None
    SimpleDocTemplate = None
    Spacer = None


def export_root() -> Path:
    return Path(
        getattr(settings, "REPORT_EXPORT_ROOT", Path(getattr(settings, "BASE_DIR", ".")).parent / "generated_reports")
    )


def ensure_export_root() -> Path:
    root = export_root()
    root.mkdir(parents=True, exist_ok=True)
    return root


def normalize_report_mode(report_mode: str | None) -> str:
    return "brief" if str(report_mode).lower() == "brief" else "full"


def get_report_content(report: Report, report_mode: str = "full") -> str:
    mode = normalize_report_mode(report_mode)
    if mode == "brief":
        if report.content_brief and report.content_brief.strip():
            return report.content_brief.strip()
        return build_brief_from_markdown(report)
    return (report.content_markdown or "").strip()


def build_brief_from_markdown(report: Report) -> str:
    markdown = (report.content_markdown or "").strip()
    if not markdown:
        return report.summary or ""

    lines = [line.strip() for line in markdown.splitlines() if line.strip()]
    brief_lines = []
    if report.summary:
        brief_lines.append(f"# {report.title}")
        brief_lines.append("")
        brief_lines.append(report.summary.strip())
        brief_lines.append("")

    heading_count = 0
    for line in lines:
        if line.startswith("#"):
            heading_count += 1
            if heading_count > 3:
                break
            brief_lines.append(line)
            continue
        if heading_count <= 3:
            brief_lines.append(line)
        if len("\n".join(brief_lines)) > 1200:
            break

    if not brief_lines:
        return markdown[:1200]
    return "\n".join(brief_lines).strip()


def build_markdown_document(report: Report, report_mode: str = "full") -> str:
    content = get_report_content(report, report_mode)
    citations = list(report.citations.order_by("index_number").values(
        "index_number", "source_title", "source_url", "cited_text_snippet"
    ))
    parts = [
        f"# {report.title}",
        "",
        f"> 摘要：{report.summary}",
        "",
        content,
    ]

    if citations:
        parts.extend(["", "## 引用来源", ""])
        for item in citations:
            snippet = f" - 引文：{item['cited_text_snippet']}" if item["cited_text_snippet"] else ""
            parts.append(
                f"{item['index_number']}. {item['source_title']} ({item['source_url']}){snippet}"
            )

    return "\n".join(part for part in parts if part is not None).strip() + "\n"


def markdown_to_plain_text(markdown_text: str) -> str:
    text = markdown_text.replace("\r\n", "\n")
    text = re.sub(r"```.*?```", "", text, flags=re.S)
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = re.sub(r"!\[[^\]]*\]\([^)]+\)", "", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"^#{1,6}\s*", "", text, flags=re.M)
    text = re.sub(r"^\>\s?", "", text, flags=re.M)
    text = re.sub(r"^\-\s+", "• ", text, flags=re.M)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def report_export_basename(report: Report, report_mode: str, extension: str) -> str:
    safe_title = re.sub(r'[\\/:*?"<>|]+', "_", report.title).strip() or f"report_{report.id}"
    return f"report_{report.id}_v{report.version}_{report_mode}_{safe_title}.{extension}"


def export_markdown(report: Report, report_mode: str = "full") -> tuple[str, str]:
    output_dir = ensure_export_root()
    filename = report_export_basename(report, report_mode, "md")
    markdown_text = build_markdown_document(report, report_mode)
    file_path = output_dir / filename
    file_path.write_text(markdown_text, encoding="utf-8")
    return str(file_path), markdown_text


def export_html(report: Report, report_mode: str = "full") -> str:
    markdown_text = build_markdown_document(report, report_mode)
    plain_text = markdown_to_plain_text(markdown_text)
    output_dir = ensure_export_root()
    filename = report_export_basename(report, report_mode, "html")
    file_path = output_dir / filename

    escaped_title = (
        report.title.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    )
    body = "".join(
        f"<p>{block.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace(chr(10), '<br/>')}</p>"
        for block in _iter_blocks(plain_text)
    )
    html = (
        "<!DOCTYPE html>"
        "<html lang=\"zh-CN\"><head><meta charset=\"utf-8\">"
        f"<title>{escaped_title}</title>"
        "<style>body{font-family:Arial,'Microsoft YaHei',sans-serif;max-width:960px;margin:40px auto;line-height:1.8;color:#1f2937;padding:0 16px;}h1{margin-bottom:24px;}p{margin:0 0 16px;white-space:pre-wrap;}</style>"
        f"</head><body><h1>{escaped_title}</h1>{body}</body></html>"
    )
    file_path.write_text(html, encoding="utf-8")
    return str(file_path)


def export_docx(report: Report, report_mode: str = "full") -> str:
    if Document is None:
        raise RuntimeError("缺少 python-docx 依赖，无法导出 Word")

    markdown_text = build_markdown_document(report, report_mode)
    plain_text = markdown_to_plain_text(markdown_text)
    output_dir = ensure_export_root()
    filename = report_export_basename(report, report_mode, "docx")
    file_path = output_dir / filename

    doc = Document()
    doc.add_heading(report.title, level=0)
    for block in _iter_blocks(plain_text):
        doc.add_paragraph(block)
    doc.save(file_path)
    return str(file_path)


def export_pdf(report: Report, report_mode: str = "full") -> str:
    if not all([
        A4, colors, getSampleStyleSheet, HRFlowable, mm, ParagraphStyle, pdfmetrics,
        TA_CENTER is not None, TA_LEFT is not None, Table, TableStyle, UnicodeCIDFont,
        Paragraph, SimpleDocTemplate, Spacer,
    ]):
        raise RuntimeError("缺少 reportlab 依赖，无法导出 PDF")

    output_dir = ensure_export_root()
    filename = report_export_basename(report, report_mode, "pdf")
    file_path = output_dir / filename

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    pdf_styles = _pdf_styles()
    story = _build_pdf_story(report, report_mode, pdf_styles)

    doc = SimpleDocTemplate(
        str(file_path),
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=22 * mm,
        bottomMargin=18 * mm,
        title=report.title,
        author="8Feet",
    )
    doc.build(
        story,
        onFirstPage=lambda canvas, document: _draw_pdf_page(canvas, document, report.title),
        onLaterPages=lambda canvas, document: _draw_pdf_page(canvas, document, report.title),
    )
    return str(file_path)


def _iter_blocks(text: str) -> Iterable[str]:
    return [block.strip() for block in text.split("\n\n") if block.strip()]


def _pdf_styles() -> dict[str, ParagraphStyle]:
    base = getSampleStyleSheet()
    body = ParagraphStyle(
        "ReportBody",
        parent=base["BodyText"],
        fontName="STSong-Light",
        fontSize=10.5,
        leading=18,
        textColor=colors.HexColor("#1F2937"),
        spaceAfter=7,
        alignment=TA_LEFT,
    )
    return {
        "title": ParagraphStyle(
            "ReportTitle",
            parent=body,
            fontSize=22,
            leading=30,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#111827"),
            spaceAfter=12,
        ),
        "subtitle": ParagraphStyle(
            "ReportSubtitle",
            parent=body,
            fontSize=9,
            leading=13,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#6B7280"),
            spaceAfter=18,
        ),
        "body": body,
        "summary": ParagraphStyle(
            "ReportSummary",
            parent=body,
            fontSize=10.5,
            leading=18,
            textColor=colors.HexColor("#374151"),
            leftIndent=2,
            rightIndent=2,
        ),
        "h1": ParagraphStyle(
            "ReportHeading1",
            parent=body,
            fontSize=15,
            leading=22,
            textColor=colors.HexColor("#111827"),
            spaceBefore=14,
            spaceAfter=8,
        ),
        "h2": ParagraphStyle(
            "ReportHeading2",
            parent=body,
            fontSize=13,
            leading=20,
            textColor=colors.HexColor("#1F2937"),
            spaceBefore=10,
            spaceAfter=6,
        ),
        "h3": ParagraphStyle(
            "ReportHeading3",
            parent=body,
            fontSize=11.5,
            leading=18,
            textColor=colors.HexColor("#374151"),
            spaceBefore=8,
            spaceAfter=4,
        ),
        "bullet": ParagraphStyle(
            "ReportBullet",
            parent=body,
            leftIndent=14,
            firstLineIndent=-9,
            bulletIndent=0,
            spaceAfter=4,
        ),
        "quote": ParagraphStyle(
            "ReportQuote",
            parent=body,
            leftIndent=10,
            rightIndent=6,
            textColor=colors.HexColor("#4B5563"),
            borderColor=colors.HexColor("#CBD5E1"),
            borderWidth=0,
            borderPadding=5,
            backColor=colors.HexColor("#F8FAFC"),
        ),
        "code": ParagraphStyle(
            "ReportCode",
            parent=body,
            fontName="STSong-Light",
            fontSize=8.5,
            leading=13,
            textColor=colors.HexColor("#111827"),
        ),
        "citation": ParagraphStyle(
            "ReportCitation",
            parent=body,
            fontSize=8.8,
            leading=14,
            textColor=colors.HexColor("#374151"),
            leftIndent=10,
            firstLineIndent=-10,
            spaceAfter=5,
        ),
        "meta": ParagraphStyle(
            "ReportMeta",
            parent=body,
            fontSize=8.5,
            leading=12,
            textColor=colors.HexColor("#374151"),
        ),
    }


def _build_pdf_story(report: Report, report_mode: str, styles: dict[str, ParagraphStyle]) -> list:
    mode_label = "简版报告" if normalize_report_mode(report_mode) == "brief" else "详版报告"
    created_at = _format_pdf_datetime(report.created_at)
    content = _strip_duplicate_title(get_report_content(report, report_mode), report.title)
    citations = list(report.citations.order_by("index_number").values(
        "index_number", "source_title", "source_url", "cited_text_snippet"
    ))

    story: list = [
        Paragraph(_pdf_inline(report.title), styles["title"]),
        Paragraph(f"8Feet 商业对象智能深度调研 · {mode_label}", styles["subtitle"]),
        _meta_table(report, mode_label, created_at, styles),
        Spacer(1, 12),
    ]

    if report.summary:
        story.extend([
            Paragraph("摘要", styles["h2"]),
            _summary_box(report.summary, styles),
            Spacer(1, 6),
        ])

    story.extend([
        HRFlowable(width="100%", thickness=0.8, color=colors.HexColor("#D1D5DB"), spaceBefore=4, spaceAfter=10),
        *_markdown_to_pdf_flowables(content, styles),
    ])

    if citations:
        story.extend([
            Spacer(1, 12),
            HRFlowable(width="100%", thickness=0.8, color=colors.HexColor("#D1D5DB"), spaceBefore=6, spaceAfter=10),
            Paragraph("引用来源", styles["h1"]),
        ])
        for item in citations:
            number = item["index_number"] or len(story)
            source_url = f"<br/><font color='#2563EB'>{_pdf_inline(item['source_url'])}</font>" if item["source_url"] else ""
            snippet = f"<br/><font color='#6B7280'>引文：{_pdf_inline(item['cited_text_snippet'])}</font>" if item["cited_text_snippet"] else ""
            story.append(Paragraph(f"{number}. {_pdf_inline(item['source_title'])}{source_url}{snippet}", styles["citation"]))

    return story


def _meta_table(report: Report, mode_label: str, created_at: str, styles: dict[str, ParagraphStyle]):
    label_style = ParagraphStyle(
        "ReportMetaLabel",
        parent=styles["meta"],
        textColor=colors.HexColor("#6B7280"),
    )
    data = [
        [
            Paragraph("报告 ID", label_style),
            Paragraph(str(report.id), styles["meta"]),
            Paragraph("任务 ID", label_style),
            Paragraph(str(report.task_id), styles["meta"]),
        ],
        [
            Paragraph("版本", label_style),
            Paragraph(f"v{report.version}", styles["meta"]),
            Paragraph("报告范围", label_style),
            Paragraph(mode_label, styles["meta"]),
        ],
        [
            Paragraph("生成时间", label_style),
            Paragraph(created_at, styles["meta"]),
            Paragraph("导出平台", label_style),
            Paragraph("8Feet", styles["meta"]),
        ],
    ]
    table = Table(data, colWidths=[22 * mm, 55 * mm, 22 * mm, 55 * mm], hAlign="CENTER")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ("BOX", (0, 0), (-1, -1), 0.6, colors.HexColor("#CBD5E1")),
        ("INNERGRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E5E7EB")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 7),
        ("RIGHTPADDING", (0, 0), (-1, -1), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    return table


def _summary_box(summary: str, styles: dict[str, ParagraphStyle]):
    table = Table([[Paragraph(_pdf_inline(summary), styles["summary"])]], colWidths=[154 * mm])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ("BOX", (0, 0), (-1, -1), 0.8, colors.HexColor("#94A3B8")),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 8),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    return table


def _markdown_to_pdf_flowables(markdown_text: str, styles: dict[str, ParagraphStyle]) -> list:
    flowables: list = []
    paragraph_lines: list[str] = []
    code_lines: list[str] = []
    in_code_block = False

    def flush_paragraph():
        if not paragraph_lines:
            return
        text = " ".join(line.strip() for line in paragraph_lines if line.strip())
        if text:
            flowables.append(Paragraph(_pdf_inline(text), styles["body"]))
        paragraph_lines.clear()

    def flush_code():
        if not code_lines:
            return
        code = "<br/>".join(escape(line) or "&nbsp;" for line in code_lines)
        table = Table([[Paragraph(code, styles["code"])]], colWidths=[154 * mm])
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F3F4F6")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#D1D5DB")),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
            ("RIGHTPADDING", (0, 0), (-1, -1), 8),
            ("TOPPADDING", (0, 0), (-1, -1), 6),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ]))
        flowables.extend([table, Spacer(1, 6)])
        code_lines.clear()

    for raw_line in markdown_text.replace("\r\n", "\n").splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code_block:
                flush_code()
                in_code_block = False
            else:
                flush_paragraph()
                in_code_block = True
            continue

        if in_code_block:
            code_lines.append(line)
            continue

        if not stripped:
            flush_paragraph()
            continue

        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            style_name = "h1" if level == 1 else "h2" if level == 2 else "h3"
            flowables.append(Paragraph(_pdf_inline(heading.group(2)), styles[style_name]))
            continue

        bullet = re.match(r"^\s*[-*+]\s+(.+)$", line)
        if bullet:
            flush_paragraph()
            flowables.append(Paragraph(_pdf_inline(bullet.group(1)), styles["bullet"], bulletText="•"))
            continue

        numbered = re.match(r"^\s*(\d+)[.)]\s+(.+)$", line)
        if numbered:
            flush_paragraph()
            flowables.append(Paragraph(_pdf_inline(numbered.group(2)), styles["bullet"], bulletText=f"{numbered.group(1)}."))
            continue

        quote = re.match(r"^>\s?(.+)$", stripped)
        if quote:
            flush_paragraph()
            flowables.append(Paragraph(_pdf_inline(quote.group(1)), styles["quote"]))
            continue

        paragraph_lines.append(line)

    flush_paragraph()
    flush_code()
    return flowables


def _strip_duplicate_title(markdown_text: str, title: str) -> str:
    pattern = rf"^\s*#\s+{re.escape(title.strip())}\s*"
    return re.sub(pattern, "", markdown_text or "", count=1).strip()


def _pdf_inline(text: object) -> str:
    value = str(text or "")
    value = re.sub(r"!\[[^\]]*\]\([^)]+\)", "", value)
    value = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", value)
    value = escape(value)
    value = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", value)
    value = re.sub(r"`([^`]+)`", r"<font color='#111827'>\1</font>", value)
    return value.replace("\n", "<br/>")


def _format_pdf_datetime(value) -> str:
    if not value:
        return ""
    if hasattr(value, "isoformat"):
        return timezone.localtime(value).strftime("%Y-%m-%d %H:%M:%S")
    return str(value)


def _draw_pdf_page(canvas, doc, title: str):
    width, height = A4
    canvas.saveState()
    canvas.setFont("STSong-Light", 8.5)
    canvas.setFillColor(colors.HexColor("#6B7280"))
    header_title = title if len(title) <= 34 else f"{title[:34]}..."
    canvas.drawString(doc.leftMargin, height - 12 * mm, header_title)
    canvas.drawRightString(width - doc.rightMargin, height - 12 * mm, "8Feet Research Report")
    canvas.setStrokeColor(colors.HexColor("#E5E7EB"))
    canvas.setLineWidth(0.4)
    canvas.line(doc.leftMargin, height - 14 * mm, width - doc.rightMargin, height - 14 * mm)
    canvas.drawCentredString(width / 2, 10 * mm, str(doc.page))
    canvas.restoreState()
